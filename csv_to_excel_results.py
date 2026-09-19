"""Convert pipeline RESULTS_*.csv files into formatted Excel result workbooks.

The generated workbook keeps the raw pipeline columns and adds a formatted
NOXIM-style summary sheet similar to ``2.5D_results (1).xlsx``.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


RAW_SHEET = "pipeline_csv"
SUMMARY_SHEET = "noxim_results"

METRIC_COLUMNS = [
    ("Global avg Delay", "Global average delay (cycles)"),
    ("Global avg Throughput", "Network throughput (flits/cycle)"),
    ("Total Energy_J", "Total energy (J)"),
    ("tot_recevd Pkts", "Total received packets"),
    ("avg Pkt Energy in uJ", None),
]

DIMENSION_HEADER = "x,y,2.5w,3h,n2.5,n3"

ALGORITHM_FILLS = {
    "GA": "C6EFCE",
    "SA": "BDD7EE",
    "PSO": "FFC7CE",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert a pipeline RESULTS_*.csv file to a formatted .xlsx workbook."
    )
    parser.add_argument("csv", type=Path, help="Input pipeline CSV, e.g. native_3d_mesh/RESULTS_elevator_first_20260919_143000.csv")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Output .xlsx path. Defaults to the CSV name with .xlsx extension.",
    )
    return parser.parse_args()


def require_columns(df: pd.DataFrame, columns: Iterable[str]) -> None:
    missing = [column for column in columns if column not in df.columns]
    if missing:
        raise ValueError(f"CSV is missing required columns: {', '.join(missing)}")


def normalize_value(value):
    if pd.isna(value):
        return None
    if hasattr(value, "item"):
        return value.item()
    return value


def validate_dimension(dimension: object) -> None:
    parts = [part.strip() for part in str(dimension).split(",")]
    if len(parts) != 6:
        raise ValueError(
            f"Dimension must have 6 values as x,y,2.5w,3h,n2.5,n3; got {dimension!r}"
        )
    for part in parts:
        int(part)


def graph_file_name(graph: object) -> str:
    text = str(graph).strip()
    match = re.search(r"Graph\s*(\d+)", text, flags=re.IGNORECASE)
    if match:
        return f"Graph{match.group(1)}.txt"
    if re.fullmatch(r"\d+", text):
        return f"Graph{text}.txt"
    return Path(text).name


def core_count(graph: object, repo_root: Path) -> int:
    graph_path = repo_root / "Graphs" / graph_file_name(graph)
    with graph_path.open("r", encoding="utf-8") as handle:
        first_line = handle.readline().strip()
    try:
        return int(first_line)
    except ValueError as exc:
        raise ValueError(f"Cannot read core count from {graph_path}") from exc


def benchmark_label(graph: object) -> str:
    text = str(graph)
    return text if text.lower().startswith("graph") else f"Graph{text}"


def metric_value(row: pd.Series, source_column: str | None):
    if source_column is None:
        packets = row.get("Total received packets")
        energy = row.get("Total energy (J)")
        if pd.isna(packets) or not packets:
            return None
        return float(energy) / float(packets) * 1_000_000
    return row.get(source_column)


def write_raw_sheet(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(RAW_SHEET)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([normalize_value(value) for value in row])
    style_table(ws)


def write_summary_sheet(wb: Workbook, df: pd.DataFrame, repo_root: Path) -> None:
    ws = wb.active
    ws.title = SUMMARY_SHEET

    groups = []
    for key, group_df in df.groupby(["Algorithm", "Population", "Iterations"], sort=True):
        groups.append((key, group_df.reset_index(drop=True)))

    ws.cell(1, 1, "Benchmark")
    ws.cell(1, 2, DIMENSION_HEADER)
    core_column = 3
    first_metric_column = core_column + 1
    ws.cell(1, core_column, "# Cores")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    for column in range(2, first_metric_column):
        ws.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)

    group_ranges = []
    start_col = first_metric_column
    for (algorithm, population, iterations), _ in groups:
        title = f"{algorithm}_{population}P_{iterations}I"
        end_col = start_col + len(METRIC_COLUMNS) - 1
        group_ranges.append((algorithm, start_col, end_col))
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        ws.cell(1, start_col, title)
        for offset, (label, _) in enumerate(METRIC_COLUMNS):
            ws.cell(2, start_col + offset, label)
        start_col = end_col + 1

    row_keys = (
        df[["Graph", "Dimension"]]
        .drop_duplicates()
        .sort_values(["Graph", "Dimension"], key=lambda col: col.astype(str))
        .itertuples(index=False, name=None)
    )

    row_number = 3
    for graph, dimension in row_keys:
        validate_dimension(dimension)
        ws.cell(row_number, 1, benchmark_label(graph))
        ws.cell(row_number, 2, dimension)
        ws.cell(row_number, core_column, core_count(graph, repo_root))

        start_col = first_metric_column
        for _, group_df in groups:
            matches = group_df[(group_df["Graph"] == graph) & (group_df["Dimension"] == dimension)]
            if not matches.empty:
                row = matches.iloc[0]
                for offset, (_, source_column) in enumerate(METRIC_COLUMNS):
                    ws.cell(row_number, start_col + offset, normalize_value(metric_value(row, source_column)))
            start_col += len(METRIC_COLUMNS)
        row_number += 1

    style_table(ws)
    color_algorithm_groups(ws, group_ranges)


def color_algorithm_groups(ws, group_ranges: list[tuple[str, int, int]]) -> None:
    for algorithm, start_col, end_col in group_ranges:
        fill_color = ALGORITHM_FILLS.get(str(algorithm).upper())
        if not fill_color:
            continue
        fill = PatternFill("solid", fgColor=fill_color)
        for row in range(1, ws.max_row + 1):
            for column in range(start_col, end_col + 1):
                ws.cell(row, column).fill = fill


def style_table(ws) -> None:
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    subheader_fill = PatternFill("solid", fgColor="EAF3F8")
    bold = Font(bold=True)

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if cell.row == 1:
                cell.font = bold
                cell.fill = header_fill
            elif cell.row == 2:
                cell.font = bold
                cell.fill = subheader_fill

    for column_cells in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 12), 28)


def convert(csv_path: Path, output_path: Path) -> None:
    df = pd.read_csv(csv_path)
    repo_root = Path(__file__).resolve().parent
    required = [
        "Graph",
        "Algorithm",
        "Dimension",
        "Population",
        "Iterations",
        "Global average delay (cycles)",
        "Network throughput (flits/cycle)",
        "Total energy (J)",
        "Total received packets",
    ]
    require_columns(df, required)

    wb = Workbook()
    write_summary_sheet(wb, df, repo_root)
    write_raw_sheet(wb, df)
    wb.save(output_path)


def main() -> None:
    args = parse_args()
    csv_path = args.csv
    output_path = args.output or csv_path.with_suffix(".xlsx")
    convert(csv_path, output_path)
    print(f"Wrote {output_path}")


if __name__ == "__main__":
    main()
