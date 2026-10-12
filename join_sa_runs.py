"""Outer join two result files by graph name.

Inputs can be formatted result workbooks produced by csv_to_excel_results.py
or raw pipeline CSV files. The output keeps the same two-row grouped result
format as the source Excel files.

Default comparison:
  - SA_Random_Run1.xlsx
  - SA_ElevatorFirst_Run1.xlsx

Example:
  python join_sa_runs.py run_a.csv run_b.csv --left-label RunA --right-label RunB
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_RANDOM = Path("SA_Random_Run1.xlsx")
DEFAULT_ELEVATOR = Path("SA_ElevatorFirst_Run1.xlsx")
DEFAULT_OUTPUT = Path("SA_Random_vs_ElevatorFirst_Run1.xlsx")
SHEET_NAME = "noxim_results"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Outer join two CSV/XLSX result files on graph name."
    )
    parser.add_argument("left", nargs="?", type=Path, default=DEFAULT_RANDOM, help="Left CSV/XLSX file.")
    parser.add_argument("right", nargs="?", type=Path, default=DEFAULT_ELEVATOR, help="Right CSV/XLSX file.")
    parser.add_argument("--left-label", default=None, help="Merged header label for the left file.")
    parser.add_argument(
        "--right-label",
        default=None,
        help="Merged header label for the right file.",
    )
    parser.add_argument("--sheet", default=SHEET_NAME, help="Sheet to read from XLSX inputs.")
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT, help="Output workbook.")
    return parser.parse_args()


def graph_sort_key(graph_name: object) -> int:
    match = re.search(r"(\d+)", str(graph_name))
    return int(match.group(1)) if match else 10**9


def default_label(path: Path) -> str:
    return path.stem


def benchmark_label(graph: object) -> str:
    text = str(graph).strip()
    return text if text.lower().startswith("graph") else f"Graph{text}"


def read_formatted_xlsx(path: Path, sheet_name: str) -> pd.DataFrame:
    raw = pd.read_excel(path, sheet_name=sheet_name, header=None)
    if raw.shape[0] < 3:
        raise ValueError(f"{path} does not look like a formatted result workbook")

    metric_header = raw.iloc[1].tolist()
    columns = []

    for index in range(raw.shape[1]):
        metric_value = metric_header[index]
        if index == 0:
            columns.append("Benchmark")
        elif index == 1:
            columns.append("Dimension")
        elif index == 2:
            columns.append("# Cores")
        else:
            columns.append(str(metric_value))

    data = raw.iloc[2:].copy()
    data.columns = columns
    data = data[data["Benchmark"].notna()]
    data = data.drop_duplicates(subset=["Benchmark"], keep="first")
    return data


def read_pipeline_csv(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path)
    required = [
        "Graph",
        "Dimension",
        "Hopcount",
        "Variance",
        "Global average delay (cycles)",
        "Network throughput (flits/cycle)",
        "Total energy (J)",
        "Total received packets",
    ]
    missing = [column for column in required if column not in df.columns]
    if missing:
        raise ValueError(f"{path} is missing required columns: {', '.join(missing)}")

    out = pd.DataFrame()
    out["Benchmark"] = df["Graph"].apply(benchmark_label)
    out["Dimension"] = df["Dimension"]
    out["# Cores"] = df["Graph"].apply(lambda graph: read_core_count(graph, path.parent))
    out["Hopcount"] = df["Hopcount"]
    out["Variance"] = df["Variance"]
    out["Global avg Delay"] = df["Global average delay (cycles)"]
    out["Global avg Throughput"] = df["Network throughput (flits/cycle)"]
    out["Total Energy_J"] = df["Total energy (J)"]
    out["tot_recevd Pkts"] = df["Total received packets"]
    out["avg Pkt Energy in uJ"] = (
        df["Total energy (J)"] / df["Total received packets"] * 1_000_000
    )
    return out.drop_duplicates(subset=["Benchmark"], keep="first")


def read_core_count(graph: object, start_dir: Path) -> int | None:
    graph_file = graph_file_name(graph)
    for root in [start_dir, *start_dir.parents]:
        candidate = root / "Graphs" / graph_file
        if candidate.exists():
            with candidate.open("r", encoding="utf-8") as handle:
                return int(handle.readline().strip())
    return None


def graph_file_name(graph: object) -> str:
    text = str(graph).strip()
    match = re.search(r"Graph\s*(\d+)", text, flags=re.IGNORECASE)
    if match:
        return f"Graph{match.group(1)}.txt"
    if re.fullmatch(r"\d+", text):
        return f"Graph{text}.txt"
    return Path(text).name


def read_result_file(path: Path, sheet_name: str) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        return read_pipeline_csv(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return read_formatted_xlsx(path, sheet_name)
    raise ValueError(f"Unsupported file type: {path}")


def join_workbooks(
    left_path: Path,
    right_path: Path,
    output_path: Path,
    left_label: str,
    right_label: str,
    sheet_name: str,
) -> None:
    left_df = read_result_file(left_path, sheet_name)
    right_df = read_result_file(right_path, sheet_name)

    joined = left_df.merge(
        right_df,
        on="Benchmark",
        how="outer",
        suffixes=("_Left", "_Right"),
    )

    joined["CoreSortKey"] = joined.apply(
        lambda row: first_non_null(row.get("# Cores_Left"), row.get("# Cores_Right")),
        axis=1,
    )
    joined["GraphSortKey"] = joined["Benchmark"].apply(graph_sort_key)
    joined = joined.sort_values(["CoreSortKey", "GraphSortKey", "Benchmark"]).drop(
        columns=["CoreSortKey", "GraphSortKey"]
    )

    write_joined_workbook(joined, output_path, left_label, right_label)


def first_non_null(left, right):
    return left if pd.notna(left) else right


def column_value(row: pd.Series, base_name: str, side: str):
    suffixed = f"{base_name}_{side}"
    if suffixed in row:
        return row[suffixed]
    return row.get(base_name)


def write_joined_workbook(
    joined: pd.DataFrame,
    output_path: Path,
    left_label: str,
    right_label: str,
) -> None:
    metric_names = [
        "Hopcount",
        "Variance",
        "Global avg Delay",
        "Global avg Throughput",
        "Total Energy_J",
        "tot_recevd Pkts",
        "avg Pkt Energy in uJ",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "joined_results"

    ws.cell(1, 1, "Benchmark")
    ws.cell(1, 2, "x,y,2.5w,3h,n2.5,n3")
    ws.cell(1, 3, "# Cores")
    for column in range(1, 4):
        ws.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)

    left_start = 4
    left_end = left_start + len(metric_names) - 1
    right_start = left_end + 1
    right_end = right_start + len(metric_names) - 1

    ws.merge_cells(start_row=1, start_column=left_start, end_row=1, end_column=left_end)
    ws.cell(1, left_start, left_label)
    ws.merge_cells(start_row=1, start_column=right_start, end_row=1, end_column=right_end)
    ws.cell(1, right_start, right_label)

    for offset, metric_name in enumerate(metric_names):
        ws.cell(2, left_start + offset, metric_name)
        ws.cell(2, right_start + offset, metric_name)

    for row_number, (_, row) in enumerate(joined.iterrows(), start=3):
        dimension = first_non_null(row.get("Dimension_Left"), row.get("Dimension_Right"))
        cores = first_non_null(row.get("# Cores_Left"), row.get("# Cores_Right"))
        ws.cell(row_number, 1, row["Benchmark"])
        ws.cell(row_number, 2, dimension)
        ws.cell(row_number, 3, cores)
        for offset, metric_name in enumerate(metric_names):
            ws.cell(row_number, left_start + offset, column_value(row, metric_name, "Left"))
            ws.cell(row_number, right_start + offset, column_value(row, metric_name, "Right"))

    style_output(ws, left_start, left_end, right_start, right_end)
    wb.save(output_path)


def style_output(ws, left_start: int, left_end: int, right_start: int, right_end: int) -> None:

    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    left_fill = PatternFill("solid", fgColor="BDD7EE")
    right_fill = PatternFill("solid", fgColor="D9EAD3")
    bold = Font(bold=True)

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if cell.row == 1:
                cell.font = bold
                cell.fill = header_fill
            elif left_start <= cell.column <= left_end:
                cell.fill = left_fill
            elif right_start <= cell.column <= right_end:
                cell.fill = right_fill

    for column_cells in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            max(max_length + 2, 12), 28
        )


def main() -> None:
    args = parse_args()
    join_workbooks(
        args.left,
        args.right,
        args.output,
        args.left_label or default_label(args.left),
        args.right_label or default_label(args.right),
        args.sheet,
    )
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
